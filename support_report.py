from __future__ import annotations

import json
import re
import shutil
import sqlite3
import subprocess
import tempfile
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas


class SupportReportError(RuntimeError):
    """Raised when the technical support PDF cannot be generated."""


def _digits(value: object) -> str:
    return re.sub(r"\D", "", str(value or ""))


def _key(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9]", "", text)


def _truncate(value: object, limit: int = 180) -> str:
    return str(value or "").strip().replace("\r", " ").replace("\n", " ")[:limit]


def _equipment_fields(asset: Mapping[str, Any], device: Mapping[str, Any]) -> tuple[str, str]:
    equipment = _truncate(asset.get("modelo_excel") or asset.get("fabricante_modelo") or device.get("model"))
    if not equipment:
        return "", ""
    parts = equipment.split(maxsplit=1)
    return parts[0], parts[1] if len(parts) == 2 else ""


def _set_cell(cell: Any, value: object, *, size: int = 8) -> None:
    cell.text = _truncate(value)
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].font.size = Pt(size)


def _set_label_value(cell: Any, label: str, value: object) -> None:
    _set_cell(cell, f"{label} {_truncate(value)}".rstrip())


def _read_inventory(inventory_path: Path) -> list[dict[str, Any]]:
    try:
        loaded = json.loads(inventory_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise SupportReportError(f"No fue posible leer el inventario para el formato: {exc}") from exc
    if not isinstance(loaded, list):
        raise SupportReportError("El inventario no tiene el formato esperado.")
    return [record for record in loaded if isinstance(record, dict)]


def _lookup_context(ticket: Mapping[str, Any], inventory_db: Path, inventory_json: Path) -> dict[str, str]:
    cedula = _digits(ticket["client_user"])
    user: dict[str, str] = {}
    device: dict[str, str] = {}
    if inventory_db.exists():
        with sqlite3.connect(inventory_db) as connection:
            connection.row_factory = sqlite3.Row
            row = connection.execute(
                "SELECT rank, full_name FROM users WHERE cedula = ?",
                (cedula,),
            ).fetchone()
            if row:
                user = dict(row)
            rows = connection.execute(
                """
                SELECT dependency, hostname, serial_1, serial_2, inventory_number, model
                FROM equipment_assignments assignment
                JOIN manual_devices device ON device.id = assignment.manual_device_id
                WHERE assignment.cedula = ?
                """,
                (cedula,),
            ).fetchall()
            asset_key = _key(ticket["asset"])
            exact = [
                row for row in rows
                if asset_key in {_key(row["hostname"]), _key(row["serial_1"]), _key(row["serial_2"])}
            ]
            if len(exact) == 1:
                device = dict(exact[0])
            elif len(rows) == 1:
                device = dict(rows[0])

    inventory = _read_inventory(inventory_json)
    asset_key = _key(ticket["asset"])
    selected = [
        item for item in inventory
        if asset_key in {
            _key(item.get("hostname")),
            _key(item.get("hostname_excel")),
            _key(item.get("hostname_pdf")),
            _key(item.get("serial_excel")),
            _key(item.get("serial_pdf")),
        }
    ]
    if len(selected) != 1 and cedula:
        selected = [
            item for item in inventory
            if cedula in {
                _digits(item.get("usuario_responsable")),
                _digits(item.get("cedula_excel")),
                _digits(item.get("usuario_actual")),
            }
        ]
    asset = selected[0] if len(selected) == 1 else {}
    brand, model = _equipment_fields(asset, device)
    return {
        "rank": _truncate(user.get("rank")),
        "full_name": _truncate(user.get("full_name") or ticket["client_name"]),
        "dependency": _truncate(asset.get("dependencia") or device.get("dependency")),
        "brand": brand,
        "model": model,
        "serial": _truncate(asset.get("serial_excel") or asset.get("serial_pdf") or device.get("serial_1")),
        "operating_system": _truncate(asset.get("so")),
        "processor": _truncate(asset.get("procesador")),
        "memory": _truncate(asset.get("ram_detalle") or (f"{asset['ram_gb']} GB" if asset.get("ram_gb") else "")),
        "disk": _truncate(", ".join(asset.get("discos") or [])),
        "ip": _truncate(asset.get("ip")),
        "hostname": _truncate(asset.get("hostname") or device.get("hostname")),
    }


def find_libreoffice() -> Path | None:
    candidates = [
        Path(r"C:\Program Files\LibreOffice\program\soffice.com"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.com"),
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    ]
    for command in ("soffice", "libreoffice"):
        located = shutil.which(command)
        if located:
            candidates.append(Path(located))
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    return None


def _convert_docx_to_pdf_libreoffice(source: Path, destination: Path) -> None:
    executable = find_libreoffice()
    if executable is None:
        raise SupportReportError("LibreOffice no está instalado o soffice.exe no está disponible.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="mesa-soporte-libreoffice-") as temp_dir:
        temp_path = Path(temp_dir)
        profile_path = temp_path / "profile"
        output_path = temp_path / "output"
        output_path.mkdir()
        completed = subprocess.run(
            [
                str(executable),
                f"-env:UserInstallation={profile_path.as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_path),
                str(source),
            ],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        generated = output_path / f"{source.stem}.pdf"
        if completed.returncode != 0 or not generated.is_file() or not generated.stat().st_size:
            details = (completed.stderr or completed.stdout or "LibreOffice no produjo el archivo PDF.").strip()
            raise SupportReportError(f"LibreOffice no pudo convertir el formato a PDF: {details}")
        shutil.copyfile(generated, destination)


def _convert_docx_to_pdf_word(source: Path, destination: Path) -> None:
    script = """
param([string]$SourcePath, [string]$DestinationPath)
$word = $null
$document = $null
try {
  $word = New-Object -ComObject Word.Application
  $word.Visible = $false
  $word.DisplayAlerts = 0
  $document = $word.Documents.Open($SourcePath, $false, $true)
  $document.ExportAsFixedFormat($DestinationPath, 17)
  if (-not (Test-Path -LiteralPath $DestinationPath)) {
    throw "Microsoft Word no produjo el archivo PDF."
  }
}
finally {
  if ($document) { $document.Close($false) }
  if ($word) { $word.Quit() }
}
"""
    with tempfile.TemporaryDirectory(prefix="mesa-soporte-pdf-") as temp_dir:
        script_path = Path(temp_dir) / "convert-to-pdf.ps1"
        script_path.write_text(script, encoding="utf-8")
        completed = subprocess.run(
            [
                "powershell.exe",
                "-NoProfile",
                "-NonInteractive",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(script_path),
                str(source),
                str(destination),
            ],
            capture_output=True,
            text=True,
            timeout=90,
            check=False,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    if completed.returncode != 0 or not destination.exists() or not destination.stat().st_size:
        details = (completed.stderr or completed.stdout or "Microsoft Word no pudo exportar el PDF.").strip()
        raise SupportReportError(f"No fue posible convertir el formato a PDF: {details}")


def _convert_docx_to_pdf(source: Path, destination: Path) -> None:
    errors: list[str] = []
    for converter in (_convert_docx_to_pdf_libreoffice, _convert_docx_to_pdf_word):
        try:
            converter(source, destination)
            return
        except SupportReportError as exc:
            errors.append(str(exc))
    raise SupportReportError(" | ".join(errors))


def _draw_label_value(pdf: canvas.Canvas, x: float, y: float, width: float, label: str, value: object) -> float:
    pdf.setStrokeColor(colors.black)
    pdf.rect(x, y - 18, width, 18)
    pdf.setFont("Helvetica-Bold", 6.5)
    pdf.drawString(x + 3, y - 7, label)
    pdf.setFont("Helvetica", 7)
    text = _truncate(value, 100)
    label_width = stringWidth(label, "Helvetica-Bold", 6.5) + 7
    pdf.drawString(min(x + label_width, x + width - 5), y - 7, text)
    return y - 18


def _draw_wrapped_box(pdf: canvas.Canvas, x: float, y: float, width: float, height: float, label: str, value: object) -> float:
    pdf.rect(x, y - height, width, height)
    pdf.setFont("Helvetica-Bold", 6.5)
    pdf.drawString(x + 3, y - 8, label)
    pdf.setFont("Helvetica", 7)
    line_y = y - 18
    words = _truncate(value, 700)
    lines: list[str] = []
    current = ""
    for word in words.split():
        candidate = f"{current} {word}".strip()
        if stringWidth(candidate, "Helvetica", 7) > width - 6:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    for line in lines[: int((height - 18) / 9)]:
        pdf.drawString(x + 3, line_y, line)
        line_y -= 9
    return y - height


def _draw_header(pdf: canvas.Canvas, title: str, page: int) -> None:
    width, height = A4
    pdf.setStrokeColor(colors.black)
    pdf.setFillColor(colors.HexColor("#E6E6E6"))
    pdf.rect(36, height - 58, width - 72, 24, fill=1)
    pdf.setFillColor(colors.black)
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawCentredString(width / 2, height - 49, title)
    pdf.setFont("Helvetica", 6.5)
    pdf.drawRightString(width - 36, 20, f"Página {page} de 2")


def _render_native_pdf(
    destination: Path,
    ticket: Mapping[str, Any],
    context: Mapping[str, str],
    requested_at: datetime,
    completed_at: datetime,
) -> None:
    pdf = canvas.Canvas(str(destination), pagesize=A4)
    width, height = A4
    left = 36
    usable = width - 72
    _draw_header(pdf, "SOLICITUD SOPORTE TÉCNICO EQUIPOS DE CÓMPUTO", 1)
    y = height - 76
    _draw_label_value(pdf, left, y, usable * 0.34, "FECHA SOLICITUD:", requested_at.strftime("%Y-%m-%d"))
    _draw_label_value(pdf, left + usable * 0.34, y, usable * 0.2, "HORA:", requested_at.strftime("%H:%M"))
    _draw_label_value(pdf, left + usable * 0.54, y, usable * 0.46, "CÓDIGO SOLICITUD No:", ticket["code"])
    y -= 29
    pdf.setFillColor(colors.HexColor("#E6E6E6"))
    pdf.rect(left, y - 16, usable, 16, fill=1)
    pdf.setFillColor(colors.black)
    pdf.setFont("Helvetica-Bold", 8)
    pdf.drawString(left + 3, y - 11, "DATOS DEL SOLICITANTE")
    y -= 20
    for label, value in (
        ("USUARIO DE RED:", ticket["client_user"]),
        ("CORREO ELECTRÓNICO:", ticket["requester_email"]),
        ("GRADO:", context["rank"]),
        ("APELLIDOS Y NOMBRES:", context["full_name"]),
        ("JEFATURA/UNIDAD:", "COLOG"),
        ("COMANDO/DEPTO.:", "BRLOG1"),
        ("DEPENDENCIA:", context["dependency"]),
        ("UBICACIÓN SITIO TRABAJO:", ticket["work_location"] or ticket["asset"]),
    ):
        y = _draw_label_value(pdf, left, y, usable, label, value) - 2
    y = _draw_label_value(pdf, left, y, usable, "SOPORTE A REALIZAR:", ticket["category"]) - 2
    y = _draw_wrapped_box(
        pdf,
        left,
        y,
        usable,
        86,
        "DESCRIPCIÓN DEL SOPORTE A REALIZAR:",
        f"{ticket['subcategory']}. {ticket['description']}",
    ) - 18
    pdf.setFont("Helvetica", 7)
    pdf.line(left + 30, y, left + usable * 0.45, y)
    pdf.line(left + usable * 0.57, y, left + usable - 30, y)
    pdf.drawCentredString(left + usable * 0.23, y - 10, "Firma y posfirma Usuario")
    pdf.drawCentredString(left + usable * 0.78, y - 10, "Firma y posfirma Funcionario Computación")
    pdf.showPage()

    _draw_header(pdf, "EVALUACIÓN DEL SOPORTE TÉCNICO EQUIPOS DE CÓMPUTO", 2)
    y = height - 76
    _draw_label_value(pdf, left, y, usable * 0.5, "FECHA SOLICITUD:", requested_at.strftime("%Y-%m-%d %H:%M"))
    _draw_label_value(pdf, left + usable * 0.5, y, usable * 0.5, "CÓDIGO SOLICITUD No:", ticket["code"])
    y -= 22
    _draw_label_value(pdf, left, y, usable * 0.5, "FECHA FINALIZACIÓN:", completed_at.strftime("%Y-%m-%d %H:%M"))
    _draw_label_value(pdf, left + usable * 0.5, y, usable * 0.5, "¿FUE RESUELTA SU SOLICITUD?:", "X  SÍ    [ ] NO")
    y -= 30
    pdf.setFillColor(colors.HexColor("#E6E6E6"))
    pdf.rect(left, y - 16, usable, 16, fill=1)
    pdf.setFillColor(colors.black)
    pdf.setFont("Helvetica-Bold", 8)
    pdf.drawString(left + 3, y - 11, "DESCRIPCIÓN DEL EQUIPO")
    y -= 20
    for label, value in (
        ("MARCA:", context["brand"]),
        ("MODELO:", context["model"]),
        ("SERIAL:", context["serial"]),
        ("SISTEMA OPERATIVO:", context["operating_system"]),
        ("PROCESADOR:", context["processor"]),
        ("MEMORIA RAM:", context["memory"]),
        ("DISCO DURO:", context["disk"]),
        ("DIRECCIÓN IP:", context["ip"]),
        ("NOMBRE PC:", context["hostname"] or ticket["asset"]),
    ):
        y = _draw_label_value(pdf, left, y, usable, label, value) - 2
    y -= 7
    pdf.setFillColor(colors.HexColor("#E6E6E6"))
    pdf.rect(left, y - 16, usable, 16, fill=1)
    pdf.setFillColor(colors.black)
    pdf.setFont("Helvetica-Bold", 8)
    pdf.drawString(left + 3, y - 11, "CALIFICACIÓN SERVICIO")
    y -= 23
    pdf.setFont("Helvetica", 7)
    pdf.drawString(left, y, "Responde las siguientes preguntas con una X.  5 Excelente · 4 Muy Bueno · 3 Bueno · 2 Regular · 1 Malo")
    y -= 14
    for question in (
        "La actitud de servicio del técnico",
        "Tiempo de la atención y solución",
        "Información del avance del servicio",
        "Calificación total del servicio",
    ):
        pdf.rect(left, y - 16, usable, 16)
        pdf.setFont("Helvetica", 7)
        pdf.drawString(left + 3, y - 10, question)
        pdf.drawRightString(left + usable - 5, y - 10, "1   2   3   4   5")
        y -= 16
    y -= 8
    _draw_wrapped_box(pdf, left, y, usable, 55, "OBSERVACIONES:", "")
    pdf.save()


def generate_support_report(
    ticket: Mapping[str, Any],
    *,
    template_path: Path,
    output_directory: Path,
    inventory_db: Path,
    inventory_json: Path,
) -> Path:
    if not template_path.exists():
        raise SupportReportError("No se encontró la plantilla formato soporte tecnico.docx.")
    context = _lookup_context(ticket, inventory_db, inventory_json)
    requested_at = datetime.fromisoformat(str(ticket["created_at"]).replace("Z", "+00:00")).astimezone()
    completed_at = datetime.now().astimezone()
    document = Document(template_path)
    request, equipment = document.tables

    _set_cell(request.cell(1, 1), requested_at.strftime("%Y-%m-%d"))
    _set_cell(request.cell(1, 6), requested_at.strftime("%H:%M"))
    _set_label_value(request.cell(1, 10), "CÓDIGO SOLICITUD No:", ticket["code"])
    _set_cell(request.cell(3, 3), ticket["client_user"])
    _set_cell(request.cell(4, 3), ticket["requester_email"])
    _set_cell(request.cell(5, 3), context["rank"])
    _set_cell(request.cell(6, 3), context["full_name"])
    _set_cell(request.cell(8, 3), "COLOG")
    _set_cell(request.cell(9, 3), "BRLOG1")
    _set_cell(request.cell(10, 3), context["dependency"])
    _set_cell(request.cell(11, 3), ticket["work_location"] or ticket["asset"])
    closure_note = (
        ticket.get("closure_note", "")
        if hasattr(ticket, "get")
        else ticket["closure_note"]
    )
    _set_cell(request.cell(4, 8), closure_note, size=7)
    if ticket["category"] in {"Hardware", "Software"}:
        _set_cell(request.cell(3, 12), "X")
    elif ticket["category"] == "Impresoras":
        _set_cell(request.cell(3, 14), "X")
    elif ticket["category"] == "Red":
        _set_cell(request.cell(3, 22), "X")

    _set_cell(request.cell(15, 2), requested_at.strftime("%Y"))
    _set_cell(request.cell(15, 4), requested_at.strftime("%m"))
    _set_cell(request.cell(15, 5), requested_at.strftime("%d"))
    _set_cell(request.cell(15, 7), requested_at.strftime("%H"))
    _set_cell(request.cell(15, 9), requested_at.strftime("%M"))
    _set_cell(request.cell(15, 17), ticket["code"])
    _set_cell(request.cell(16, 2), completed_at.strftime("%Y"))
    _set_cell(request.cell(16, 4), completed_at.strftime("%m"))
    _set_cell(request.cell(16, 5), completed_at.strftime("%d"))
    _set_cell(request.cell(16, 7), completed_at.strftime("%H"))
    _set_cell(request.cell(16, 9), completed_at.strftime("%M"))
    _set_cell(request.cell(17, 17), "X")

    _set_cell(equipment.cell(2, 0), context["brand"])
    _set_cell(equipment.cell(2, 1), context["model"])
    _set_cell(equipment.cell(2, 2), context["serial"])
    _set_cell(equipment.cell(2, 5), context["operating_system"])
    _set_cell(equipment.cell(2, 6), context["processor"])
    _set_cell(equipment.cell(2, 9), context["memory"])
    _set_cell(equipment.cell(2, 10), context["disk"])
    _set_cell(equipment.cell(3, 1), context["ip"])
    _set_cell(equipment.cell(3, 6), context["hostname"] or ticket["asset"])

    output_directory.mkdir(parents=True, exist_ok=True)
    safe_code = re.sub(r"[^A-Za-z0-9_-]", "_", str(ticket["code"]))
    with tempfile.TemporaryDirectory(prefix="mesa-soporte-docx-", dir=output_directory) as temp_dir:
        docx_path = Path(temp_dir) / f"formato_soporte_{safe_code}.docx"
        pdf_path = output_directory / f"formato_soporte_{safe_code}.pdf"
        document.save(docx_path)
        try:
            _convert_docx_to_pdf(docx_path, pdf_path)
        except SupportReportError:
            _render_native_pdf(pdf_path, ticket, context, requested_at, completed_at)
    return pdf_path
